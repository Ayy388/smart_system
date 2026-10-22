# -*- coding: utf-8 -*-
"""完成实验报告5（统计组）：填写表格、插入截图、补充布局说明、问题与解决、实验总结"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "智能系统开发综合实训-实验报告5.docx"
DST = "智能系统开发综合实训-实验报告5（第91组-已完成）.docx"
SHOTS = "screenshots"

doc = docx.Document(SRC)
log = []


def set_cell(cell, text, keep_font=True, font_size=None):
    """清空单元格文字并写入 text（保留段落字体设置，支持\n换行）"""
    p = cell.paragraphs[0]
    # 记录原字体
    old_size = None
    old_bold = None
    if p.runs:
        old_size = p.runs[0].font.size
        old_bold = p.runs[0].font.bold
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    parts = text.split("\n")
    for k, part in enumerate(parts):
        run = p.add_run(part)
        sz = font_size or old_size
        if sz:
            run.font.size = sz
        if old_bold is not None:
            run.font.bold = old_bold
        if k < len(parts) - 1:
            run.add_break()
    return run


def mark_passed(cell, note=None):
    """把 '□通过 □未通过' 或 '□' 置为通过"""
    txt = cell.text.strip()
    if "□未通过" in txt or "□通过" in txt:
        set_cell(cell, ("√通过" if note is None else f"√通过（{note}）"))
    elif txt == "□":
        set_cell(cell, "√")
    elif txt == "":
        set_cell(cell, "√")
    else:
        set_cell(cell, "√")


def insert_shot(table, img_path, width_in):
    """在 1×1 截图表中插入图片并居中"""
    cell = table.rows[0].cells[0]
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    # 底部留一个空段，避免图片紧贴表格边框
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER


T = doc.tables

# ========== 1. 封面表 T0 ==========
# 组号 -> 第91组
set_cell(T[0].rows[7].cells[1], "第91组")
# 所选课题：勾选 课题C
c = T[0].rows[10].cells[1]
newtext = c.text.replace("□ 课题C：智能人流统计与考勤系统", "■ 课题C：智能人流统计与考勤系统")
set_cell(c, newtext)
log.append("封面：组号=第91组，勾选课题C")

# ========== 2. 任务总览 T2 完成打勾 ==========
for r in range(1, 9):
    mark_passed(T[2].rows[r].cells[3])
log.append("任务总览T2：8项全部打勾")

# ========== 3. 5.1 运行环境检查清单 T3 ==========
for r in range(1, 11):
    note = "vision_module.py" if r == 9 else None
    mark_passed(T[3].rows[r].cells[4], note)
log.append("5.1环境清单T3：10项通过")

# ========== 4. 5.5 前端运行条件检查 T7 ==========
set_cell(T[7].rows[1].cells[4], "√（本地echarts）")
for r in range(2, 5):
    mark_passed(T[7].rows[r].cells[4])
log.append("5.5前端条件T7：CDN改用本地，其余通过")

# ========== 5. 5.7 环境检查记录表 T9 ==========
for r in range(1, 14):
    note = None
    if r == 9:
        note = "vision_module.py"
    elif r == 12:
        note = "本地echarts"
    mark_passed(T[9].rows[r].cells[2], note)
log.append("5.7环境记录表T9：13项")

# ========== 6. 截图插入 ==========
insert_shot(T[10], f"{SHOTS}/shot0_env_check.png", 5.6)        # 截图0
insert_shot(T[13], f"{SHOTS}/shot1_terminal_initial.png", 4.4)  # 截图1
insert_shot(T[14], f"{SHOTS}/shot2_dashboard_initial.png", 5.9) # 截图2
insert_shot(T[17], f"{SHOTS}/shot3_terminal_running.png", 4.4)  # 截图3
insert_shot(T[18], f"{SHOTS}/shot4_dashboard_running.png", 5.9) # 截图4
insert_shot(T[20], f"{SHOTS}/shot5_terminal_landscape.png", 4.4)# 截图5
insert_shot(T[21], f"{SHOTS}/shot6_terminal_pc.png", 5.9)       # 截图6
log.append("插入截图0~6")

# ========== 7. 8.4.2 联调 T15 / T16 ==========
for r in range(1, 3):
    cell = T[15].rows[r].cells[1]
    set_cell(cell, cell.text + "（√通过）")
for r in range(1, 13):
    mark_passed(T[16].rows[r].cells[3])
log.append("8.4.2联调记录T15/T16")

# ========== 8. 8.7 检查表 T23 ==========
for r in range(1, 5):
    mark_passed(T[23].rows[r].cells[2])
# r5 代码已提交至Gitee：留空，需用户自行推送
log.append("8.7检查表T23：前4项√，第5项(提交Gitee)留待用户")

# ========== 9. 9.1 环境检查记录表 T24 ==========
for r in range(1, 14):
    note = None
    if r == 9:
        note = "vision_module.py"
    elif r == 12:
        note = "本地echarts"
    mark_passed(T[24].rows[r].cells[2], note)
log.append("9.1环境记录表T24：13项")

# ========== 10. 9.2 基础任务验收表 T25 ==========
for r in range(1, 12):
    mark_passed(T[25].rows[r].cells[2])
log.append("9.2基础任务验收表T25：11项")

# ========== 11. 9.3 前后端联调测试记录表 T26 ==========
for r in range(1, 11):
    mark_passed(T[26].rows[r].cells[2])
log.append("9.3联调测试记录表T26：10项")

# ========== 12. 问题与解决 T27 ==========
issues = [
    ["ECharts图表无法显示", "管理后台饼图/柱状图区域空白",
     "运行环境无法访问cdn.jsdelivr.net，外部CDN加载失败",
     "通过npm下载echarts.min.js到flask_app/static/js，HTML改为本地引用/static/js/echarts.min.js"],
    ["视频流区域全黑无画面", "/video_feed无内容",
     "运行环境缺少可用的摄像头设备，无法采集真实画面",
     "camera_manager.py增加模拟画面源，SIM=1时生成带移动人形的模拟视频，用于无摄像头联调演示"],
    ["模拟模式下计数始终为0", "进入/离开人数一直为0",
     "背景差分形态学操作内核(5×5)过大，把较小的运动目标腐蚀掉了",
     "形态学内核由5×5调整为3×3，并下调最小面积阈值，目标检测恢复正常"],
    ["页面数据显示为--", "当前/进入/离开不更新",
     "前端取值字段与后端API返回的JSON字段名不一致",
     "用curl查看API实际返回字段，与前端getElementById取值逐一比对后对齐字段名"],
]
# 清空已有空行并填充
base = T[27]
# 删除原有空数据行（保留表头）
while len(base.rows) > 1:
    tr = base.rows[1]._tr
    tr.getparent().remove(tr)
for i, row_vals in enumerate(issues):
    cells = base.add_row().cells
    for j, val in enumerate(row_vals):
        set_cell(cells[j], val)
log.append("问题与解决T27：4条")

# ========== 13. 实验总结 T28 ==========
summary = (
    "本次实验（实验五）完成了智能人流统计与考勤系统的Web前端开发，包括适配3.5英寸横屏的终端控制界面index.html和"
    "管理后台dashboard.html，并与Flask后端完成了前后端联调。通过本次实验，我对Web前端开发有了系统而深入的理解。"
    "在HTML/CSS方面，我掌握了页面结构搭建与样式设计，使用Flexbox布局实现灵活排版，并利用CSS媒体查询实现从"
    "3.5英寸横屏（480×320）到PC显示器（1920×1080）的自适应布局，深刻体会到“移动优先、响应式适配”的设计思路；"
    "同时针对小屏幕对视频比例、按钮尺寸和字号进行了专门优化。在JavaScript方面，我掌握了DOM操作、事件处理等基础，"
    "并学会使用fetch()发起AJAX异步请求：页面每2秒自动轮询/api/status、/api/latest等接口，在不刷新页面的情况下实时更新"
    "FPS、帧数、当前人数等数据，实现了流畅的实时交互，真正理解了“异步、无刷新”的数据交互方式。在ECharts方面，"
    "我使用饼图和柱状图分别展示进出人数分布与进入/离开对比，掌握了setOption配置图表数据、颜色、坐标轴等方法，"
    "体会到数据可视化能把枯燥的统计数字转化为直观的图形。在前后端联调方面，我认识到前后端字段名必须严格一致，"
    "当统计数据不显示时，通过curl检查API返回的JSON字段并与前端取值逻辑逐一比对，最终解决了字段不匹配问题；"
    "当ECharts CDN不可达导致图表无法加载时，我将图表库改为本地引用予以解决。本次实验后端还实现了基于背景差分与"
    "质心跟踪的人体检测、计数线跨越判定，以及进入/离开事件入库，管理后台可查询历史记录并生成统计图表，使系统初步闭环。"
    "在第七次课系统集成阶段，我计划将人流计数与考勤功能深度结合：基于进入/离开记录自动生成考勤打卡、统计异常缺勤，"
    "并完善实时告警、报表导出与多摄像头扩展等功能，最终形成功能完整、可直接部署的智能人流统计与考勤系统。"
)
set_cell(T[28].rows[0].cells[0], summary, font_size=Pt(10.5))
log.append(f"实验总结T28：{len(summary)}字")

# ========== 14. 7.3 统计组布局说明文字 ==========
paras = doc.paragraphs
idx_73 = next(i for i, p in enumerate(paras) if p.text.strip().startswith("7.3 统计组"))
idx_term = next(i for i in range(idx_73, len(paras))
                if paras[i].text.strip().startswith("终端控制界面（3.5英寸横屏480×320）"))
idx_dash = next(i for i in range(idx_73, len(paras))
                if paras[i].text.strip().startswith("管理后台布局"))
idx_74 = next(i for i in range(idx_73, len(paras))
              if paras[i].text.strip().startswith("7.4 三个项目界面差异对比总览"))
term_layout = ("标题栏：📊 人流统计（蓝色）+ 状态指示 + “后台”入口；视频区：实时人体检测画面与黄色计数线；"
               "结果栏：当前人数（蓝）、进入（绿）、离开（橙）；控制区：▶启动（蓝）、⏸停止（红）、🔄重置（橙）；"
               "状态栏：FPS、帧数、运行时间。")
dash_layout = ("顶部标题栏（人流数据分析 + 返回终端控制）；四张统计卡片：进入总数、离开总数、当前人数、今日统计；"
               "图表区：ECharts饼图（进出人数分布）与柱状图（进入vs离开）；下方为历史统计记录表（ID/时间/类型/人数）；"
               "底部系统设置：计数线位置、匹配距离，可保存至数据库。")
paras[idx_dash].insert_paragraph_before(term_layout)
paras[idx_74].insert_paragraph_before(dash_layout)
log.append("7.3统计组布局说明已补充")

doc.save(DST)
print("已保存:", DST)
for l in log:
    print(" -", l)
